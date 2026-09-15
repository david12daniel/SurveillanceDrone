#!/usr/bin/env python3
"""
apply_slide14_diagram.py

Insert a SysML General-View diagram (PNG) into CDR slide 14, replacing the
33 hand-drawn shapes (IDs 9-41) that make up the old system-composition figure.

Usage:
    python presentation/apply_slide14_diagram.py [--image PATH]

Defaults to the pre-generated .openclaw/tmp/aos_syson.png.
Writes the modified deck to CDR_Thermal_Surveillance_Drone.pptx in-place
(a backup is saved as *.bak.pptx first).

Based on the verified approach from apply_row0b_export.py (2026-08-17).
"""

import argparse, os, shutil, sys
from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from lxml import etree

PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(PROJ, "CDR_Thermal_Surveillance_Drone.pptx")
DEFAULT_IMAGE = os.path.join(PROJ, ".openclaw", "tmp", "aos_syson.png")

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}

# Shapes to remove (the 33 hand-drawn elements on slide 14)
SHAPE_IDS_TO_REMOVE = set(range(9, 42))  # 9 through 41 inclusive

# Picture placement — aspect-ratio preserving.
# aos_syson.png is 354x493 px (portrait); display at ~5" tall = ~3.6" wide,
# centered on the wide 13.3" slide below the title (top ~1.2", bottom ~6.9").
PICTURE_HEIGHT = Emu(int(4.9 * 914400))   # 4.9" tall
PICTURE_WIDTH  = Emu(int(4.9 * 914400 * 354.0 / 493.0))  # keep 354:493 ratio -> ~3.52"
PICTURE_LEFT   = Emu(int((13.333 * 914400 - int(4.9 * 914400 * 354.0 / 493.0)) / 2))
PICTURE_TOP    = Emu(int(1.1 * 914400))   # below header

# Fallback cross-check for non-354x493 inputs: recompute from actual image size
# (done in main() if a different image is passed).


def backup(path):
    bak = path + ".bak.pptx"
    shutil.copy2(path, bak)
    print(f"Backup saved: {bak}")


def remove_shapes(slide, ids_to_remove):
    """Remove shapes with given IDs from the slide's spTree."""
    sp_tree = slide._element.find('.//p:cSld/p:spTree', NS)
    if sp_tree is None:
        print("ERROR: cannot find spTree", file=sys.stderr)
        return False

    removed = 0
    # Iterate a copy of children list because we mutate the tree
    for child in list(sp_tree):
        # Check if this element has an nvSpPr with a matching id
        nv_sp_pr = child.find('p:nvSpPr/p:cNvPr', NS)
        if nv_sp_pr is not None:
            shape_id = int(nv_sp_pr.get('id', -1))
            if shape_id in ids_to_remove:
                sp_tree.remove(child)
                removed += 1

    print(f"Removed {removed} shapes (targeted IDs {sorted(ids_to_remove)})")
    return removed > 0


def add_picture_to_slide(slide, image_path, left, top, width, height):
    """Add a picture to the slide at the given position/size."""
    # python-pptx's add_picture adds to the presentation-level image parts
    pic = slide.shapes.add_picture(image_path, left, top, width, height)
    print(f"Added picture: {os.path.basename(image_path)} at ({left},{top}) size ({width}×{height})")
    return pic


def main():
    parser = argparse.ArgumentParser(description="Insert diagram into CDR slide 14")
    parser.add_argument("--image", default=DEFAULT_IMAGE,
                        help="PNG file to insert (default: .openclaw/tmp/aos_syson.png)")
    parser.add_argument("--deck", default=DECK,
                        help="Path to the CDR deck (default: CDR_Thermal_Surveillance_Drone.pptx)")
    parser.add_argument("--dry-run", action="store_true",
                        help="Print what would be done without modifying the file")
    args = parser.parse_args()

    if not os.path.exists(args.deck):
        print(f"ERROR: deck not found at {args.deck}", file=sys.stderr)
        sys.exit(1)
    if not os.path.exists(args.image):
        print(f"ERROR: image not found at {args.image}", file=sys.stderr)
        sys.exit(1)

    if args.dry_run:
        print(f"[DRY RUN] Would modify: {args.deck}")
        print(f"[DRY RUN] Would insert: {args.image}")
        print(f"[DRY RUN] Would remove shapes: {sorted(SHAPE_IDS_TO_REMOVE)}")
        print(f"[DRY RUN] Picture position: {PICTURE_LEFT},{PICTURE_TOP} size {PICTURE_WIDTH}×{PICTURE_HEIGHT}")
        return

    # Backup
    backup(args.deck)

    # Open deck
    prs = Presentation(args.deck)
    slide = prs.slides[13]  # slide 14 (0-indexed)

    # Verify slide 14 has expected shapes
    shape_ids_present = set()
    for s in slide.shapes:
        nv = s._element.find('p:nvSpPr/p:cNvPr', NS)
        if nv is not None:
            shape_ids_present.add(int(nv.get('id', -1)))

    target_present = SHAPE_IDS_TO_REMOVE.issubset(shape_ids_present)
    if not target_present:
        missing = SHAPE_IDS_TO_REMOVE - shape_ids_present
        print(f"WARNING: {len(missing)} of 33 target shapes not found on slide 14: {sorted(missing)}")
        print("Proceeding anyway...")
    else:
        print(f"All 33 target shapes confirmed present on slide 14.")

    # Remove old shapes
    remove_shapes(slide, SHAPE_IDS_TO_REMOVE)

    # If a different image was passed, recompute placement from its real size
    from PIL import Image
    with Image.open(args.image) as im:
        w_px, h_px = im.size
    target_h = 4.9 * 914400  # EMU
    width_emu = int(target_h * w_px / h_px)
    height_emu = int(target_h)
    left_emu = int((13.333 * 914400 - width_emu) / 2)
    top_emu = int(1.1 * 914400)

    # Add the picture
    add_picture_to_slide(slide, args.image, left_emu, top_emu,
                         width_emu, height_emu)

    # Save
    prs.save(args.deck)
    print(f"Deck saved: {args.deck}")

    # Quick verification
    prs2 = Presentation(args.deck)
    slide2 = prs2.slides[13]
    remaining = sum(1 for s in slide2.shapes
                    if s._element.find('p:nvSpPr/p:cNvPr', NS) is not None
                    and int(s._element.find('p:nvSpPr/p:cNvPr', NS).get('id', -1)) in SHAPE_IDS_TO_REMOVE)
    pictures = sum(1 for s in slide2.shapes if s.shape_type == 13)  # Picture type
    print(f"Verification: {remaining} target shapes remain (should be 0), {pictures} pictures on slide")
    print("Done.")


if __name__ == "__main__":
    main()